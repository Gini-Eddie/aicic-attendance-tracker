import uuid
import zipfile
import csv
import io
import re
from datetime import datetime
from pathlib import Path
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import Response
from pydantic import BaseModel, Field, EmailStr, TypeAdapter
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from ..database import get_db
from ..auth import require_admin, require_teacher
from ..permissions import require_course_access
from ..registration import attendance_roster, roster, cohort_key, current_cohort, resolve_student, session_cohort, lock_registration_writes
from typing import Optional
from ..models import (User, Student, Course, Enrollment, TeacherCourse, AttendanceSession, AttendanceRecord,
                      StudentRegistration, DeletedRegistration, PendingAttendance, AdminNotification,
                      NotificationRead, DeletedUser, SessionCohort, CourseCohort)

router = APIRouter(tags=["Cohort operations"])


def csv_response(filename, headers, rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    def safe(value):
        value = "" if value is None else str(value)
        return "'" + value if (value.startswith(("\t", "\r")) or value.lstrip().startswith(("=", "+", "-", "@"))) else value
    writer.writerow(headers)
    writer.writerows([[safe(value) for value in row] for row in rows])
    return Response("\ufeff" + output.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})


@router.post("/courses/{course_id}/import")
def import_students(course_id: str, file: UploadFile = File(...), cohort: str = Form(...), track: str = Form(...),
                          user: User = Depends(require_admin), db: Session = Depends(get_db)):
    require_course_access(db, user, course_id)
    cohort = cohort_key(cohort)
    if not cohort or len(cohort) > 255:
        raise HTTPException(400, "Enter a cohort name up to 255 characters.")
    slug = lambda value: re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")
    prefix = f"{slug(cohort)}-{slug(track)}"
    if not slug(track) or not slug(cohort) or len(prefix) > 52:
        raise HTTPException(400, "Enter a short cohort and track label, e.g. MATRIX and UI.")
    content = file.file.read(2 * 1024 * 1024 + 1)
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(413, "Upload a file smaller than 2 MB.")
    extension = Path(file.filename or "").suffix.lower()
    try:
        if extension == ".csv":
            reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            header = next(reader, [])
            values = []
            for row in reader:
                values.append(row)
                if len(values) > 1000:
                    raise HTTPException(400, "Import at most 1,000 rows at a time.")
        elif extension == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise HTTPException(413, "Excel workbook expands beyond 20 MB. Split it into smaller files.")
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                reader = workbook.active.iter_rows(values_only=True)
                header = next(reader, [])
                values = []
                for row in reader:
                    values.append(row)
                    if len(values) > 1000:
                        raise HTTPException(400, "Import at most 1,000 rows at a time.")
            finally:
                workbook.close()
        else:
            raise HTTPException(400, "Upload a .csv or .xlsx file with name and email columns.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Could not read the file. Use a valid CSV or Excel (.xlsx) workbook.")
    keys = [str(value or "").strip().lower().replace(" ", "_") for value in header]
    name_key = next((key for key in ("name", "full_name", "student_name") if key in keys), None)
    email_key = next((key for key in ("email", "email_address") if key in keys), None)
    if not name_key or not email_key:
        raise HTTPException(400, "The first row must contain name and email column headings.")
    # Serialize imports across workers so shared prefixes allocate numbers safely.
    lock_registration_writes(db)
    try:
        aliases = {row[0] for row in db.query(StudentRegistration.code).all()}
        used_codes = {row[0] for row in db.query(Student.student_code).all()} | aliases
        suffixes = [int(code[len(prefix)+1:]) for code in used_codes if code.startswith(prefix + "-") and code[len(prefix)+1:].isdigit()]
        number = max(suffixes, default=0) + 1
        report = {"created": [], "skipped": [], "errors": [], "cohort": cohort, "prefix": prefix}
        valid_rows = []
        for line, row in enumerate(values, 2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            try:
                name = str(row[keys.index(name_key)] or "").strip()
                email = str(TypeAdapter(EmailStr).validate_python(str(row[keys.index(email_key)] or "").strip())).lower()
                if len(name) < 2 or len(name) > 255:
                    raise ValueError("name")
                valid_rows.append((line, name, email))
            except Exception:
                report["errors"].append({"row": line, "reason": "Provide a full name (2–255 characters) and a valid email."})
        legacy_course_cohort = current_cohort(db, course_id) or cohort
        students = {r.email: r for r in db.query(Student).filter(Student.email.in_([r[2] for r in valid_rows])).all()}
        memberships = {(r.student_id, r.cohort) for r in db.query(StudentRegistration).filter_by(course_id=course_id).all()}
        registered_ids = {student_id for student_id, _ in memberships}
        enrolled_ids = {row[0] for row in db.query(Enrollment.student_id).filter_by(course_id=course_id)}
        def allocate():
            nonlocal number
            code = f"{prefix}-{number:03d}"
            while code in used_codes:
                number += 1
                code = f"{prefix}-{number:03d}"
            number += 1
            used_codes.add(code)
            return code
        for line, name, email in valid_rows:
            student = students.get(email)
            if student and (student.id, cohort) in memberships:
                report["skipped"].append({"row": line, "email": email, "reason": "Already registered in this course and cohort (including removed registrations)."})
                continue
            if student and student.id in enrolled_ids and student.id not in registered_ids:
                legacy_cohort = legacy_course_cohort
                # Preserve legacy numbers wherever they are not already an alias in another course.
                if student.student_code not in aliases:
                    db.add(StudentRegistration(student_id=student.id, course_id=course_id, cohort=legacy_cohort, code=student.student_code))
                    aliases.add(student.student_code)
                    memberships.add((student.id, legacy_cohort))
                    registered_ids.add(student.id)
                if legacy_cohort == cohort:
                    report["skipped"].append({"row": line, "email": email, "reason": "Already enrolled in this course and cohort."})
                    continue
            code = allocate()
            if not student:
                student = Student(id=str(uuid.uuid4()), full_name=name, email=email, student_code=code)
                db.add(student)
                students[email] = student
            if student.id not in enrolled_ids:
                db.add(Enrollment(student_id=student.id, course_id=course_id))
                enrolled_ids.add(student.id)
            db.add(StudentRegistration(student_id=student.id, course_id=course_id, cohort=cohort, code=code))
            aliases.add(code)
            memberships.add((student.id, cohort))
            registered_ids.add(student.id)
            report["created"].append({"name": student.full_name, "email": email, "student_code": code})
        if report["created"] and not db.get(CourseCohort, course_id):
            db.add(CourseCohort(course_id=course_id, name=cohort))
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Registrations changed concurrently. Retry the import; no changes from this upload were saved.")
    return report


@router.get("/courses/{course_id}/registrations.csv")
def export_roster(course_id: str, cohort: Optional[str] = None, user: User = Depends(require_teacher), db: Session = Depends(get_db)):
    require_course_access(db, user, course_id)
    return csv_response("students.csv", ["name", "email", "registration_number", "cohort"],
        [(item["student"].full_name, item["student"].email if not item["student"].email.endswith("@students.invalid") else "", item["code"], item["cohort"])
         for item in roster(db, course_id, cohort)])


@router.delete("/courses/{course_id}/roster/{student_id}")
def remove_student(course_id: str, student_id: str, cohort: str = "", user: User = Depends(require_teacher), db: Session = Depends(get_db)):
    require_course_access(db, user, course_id)
    student = db.get(Student, student_id)
    if not student:
        raise HTTPException(404, "Student not found")
    selected = cohort_key(cohort) or current_cohort(db, course_id)
    registrations = db.query(StudentRegistration).filter(StudentRegistration.student_id == student_id, StudentRegistration.course_id == course_id).all()
    targets = [r for r in registrations if (not selected or r.cohort == selected) and not db.get(DeletedRegistration, r.id)]
    enrollment = db.query(Enrollment).filter(Enrollment.student_id == student_id, Enrollment.course_id == course_id).first()
    if registrations and not targets or not registrations and not enrollment:
        raise HTTPException(404, "Student registration not found in this cohort")
    for item in targets:
        db.add(DeletedRegistration(registration_id=item.id))
    db.flush()
    active = [r for r in registrations if not db.get(DeletedRegistration, r.id)]
    if enrollment and not active:
        db.delete(enrollment)
    course = db.get(Course, course_id)
    db.add(AdminNotification(message=f"{user.name} removed {student.full_name} ({student.email}) from {course.name}, cohort {selected or 'legacy'}. Historical attendance was retained."))
    db.commit()
    return {"message": "Student removed from this roster. Administrators have been notified; historical attendance is retained."}


@router.get("/staff")
def staff(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    disabled = db.query(DeletedUser.user_id)
    return [{"id": item.id, "name": item.name, "email": item.email, "role": item.role,
             "courses": [{"id": c.id, "name": c.name} for c in item.courses_taught]}
            for item in db.query(User).filter(User.id.notin_(disabled)).order_by(User.name).all()]


@router.get("/staff/export.csv")
def export_staff(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    return csv_response("staff.csv", ["name", "email", "role", "courses"],
        [(item["name"], item["email"], item["role"], "; ".join(c["name"] for c in item["courses"])) for item in staff(user, db)])


@router.delete("/teachers/{teacher_id}")
def delete_teacher(teacher_id: str, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    teacher = db.get(User, teacher_id)
    if not teacher or db.get(DeletedUser, teacher_id):
        raise HTTPException(404, "Teacher not found")
    if teacher.role != "teacher":
        raise HTTPException(403, "Administrator accounts cannot be deleted through teacher management.")
    db.add(DeletedUser(user_id=teacher_id, deleted_by=user.id))
    db.query(TeacherCourse).filter(TeacherCourse.teacher_id == teacher_id).delete()
    db.query(AttendanceSession).filter(AttendanceSession.teacher_id == teacher_id, AttendanceSession.status == "active").update({"status": "closed"})
    db.add(AdminNotification(message=f"{user.name} deleted teacher access for {teacher.name}. Historical attendance and session ownership were retained."))
    db.commit()
    return {"message": "Teacher access deleted; historical records retained."}


@router.get("/notifications")
def notifications(user: User = Depends(require_admin), db: Session = Depends(get_db)):
    reads = {r.notification_id for r in db.query(NotificationRead).filter(NotificationRead.admin_id == user.id).all()}
    return [{"id": n.id, "message": n.message, "created_at": n.created_at, "read": n.id in reads}
            for n in db.query(AdminNotification).order_by(AdminNotification.created_at.desc()).limit(100).all()]


@router.post("/notifications/{notification_id}/read")
def mark_read(notification_id: str, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    if not db.get(AdminNotification, notification_id):
        raise HTTPException(404, "Notification not found")
    if not db.get(NotificationRead, (notification_id, user.id)):
        db.add(NotificationRead(notification_id=notification_id, admin_id=user.id))
        db.commit()
    return {"message": "Read"}


class DeskRequest(BaseModel):
    session_id: str
    student_code: str = Field(min_length=2, max_length=64)


@router.post("/desk/check-in", status_code=201)
def desk_checkin(data: DeskRequest, user: User = Depends(require_admin), db: Session = Depends(get_db)):
    session = db.query(AttendanceSession).filter_by(id=data.session_id).with_for_update().first()
    if not session or session.status != "active" or session.expires_at < datetime.utcnow():
        raise HTTPException(400, "Choose an active, unexpired attendance session.")
    student, code = resolve_student(db, session, data.student_code)
    if db.query(AttendanceRecord).filter(AttendanceRecord.session_id == session.id, AttendanceRecord.student_id == student.id).first():
        raise HTTPException(409, "Attendance is already recorded.")
    if db.query(PendingAttendance).filter(PendingAttendance.session_id == session.id, PendingAttendance.student_id == student.id).first():
        raise HTTPException(409, "A desk submission already exists for this student and session.")
    pending = PendingAttendance(session_id=session.id, student_id=student.id, submitted_by=user.id)
    db.add(pending)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "A desk submission already exists.")
    return {"message": "Saved as unverified. The course tutor must approve it.", "student_name": student.full_name, "student_code": code}


@router.get("/pending-attendance")
def pending_list(user: User = Depends(require_teacher), db: Session = Depends(get_db)):
    query = db.query(PendingAttendance).filter(PendingAttendance.status == "pending")
    if user.role != "admin":
        course_ids = db.query(TeacherCourse.course_id).filter(TeacherCourse.teacher_id == user.id)
        session_ids = db.query(AttendanceSession.id).filter(AttendanceSession.course_id.in_(course_ids))
        query = query.filter(PendingAttendance.session_id.in_(session_ids))
    output = []
    for item in query.order_by(PendingAttendance.submitted_at).all():
        session, student = db.get(AttendanceSession, item.session_id), db.get(Student, item.student_id)
        if not session or not student:
            continue
        can_verify = bool(db.get(TeacherCourse, (user.id, session.course_id)))
        output.append({"id": item.id, "student_name": student.full_name, "student_code": next((r["code"] for r in roster(db, session.course_id, session_cohort(db, session)) if r["student"].id == student.id), student.student_code),
                       "course_name": session.course.name, "cohort": session_cohort(db, session), "session_id": session.id,
                       "submitted_at": item.submitted_at, "source": "Admin/company laptop", "can_verify": can_verify})
    return output


class VerifyRequest(BaseModel):
    approve: bool


@router.post("/pending-attendance/{pending_id}/verify")
def verify_desk(pending_id: str, data: VerifyRequest, user: User = Depends(require_teacher), db: Session = Depends(get_db)):
    item = db.query(PendingAttendance).filter(PendingAttendance.id == pending_id).with_for_update().first()
    if not item:
        raise HTTPException(404, "Submission not found")
    session = db.get(AttendanceSession, item.session_id)
    if not session or not db.get(TeacherCourse, (user.id, session.course_id)):
        raise HTTPException(403, "Only a tutor assigned to this course can verify desk attendance.")
    if item.status != "pending":
        raise HTTPException(409, "This submission was already reviewed.")
    if data.approve:
        eligible = any(r["student"].id == item.student_id for r in roster(db, session.course_id, session_cohort(db, session)))
        if not eligible:
            raise HTTPException(409, "Student is no longer enrolled in this cohort. Reject this submission.")
        if not db.query(AttendanceRecord).filter(AttendanceRecord.session_id == session.id, AttendanceRecord.student_id == item.student_id).first():
            db.add(AttendanceRecord(session_id=session.id, student_id=item.student_id, checked_in_at=item.submitted_at))
    item.status = "verified" if data.approve else "rejected"
    item.verified_by, item.verified_at = user.id, datetime.utcnow()
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Attendance changed concurrently. Refresh and retry.")
    return {"message": "Attendance verified and recorded." if data.approve else "Submission rejected."}


@router.get("/sessions/{session_id}/export.csv")
def export_attendance(session_id: str, user: User = Depends(require_teacher), db: Session = Depends(get_db)):
    session = db.get(AttendanceSession, session_id)
    if not session:
        raise HTTPException(404, "Session not found")
    require_course_access(db, user, session.course_id)
    records = {r.student_id: r for r in db.query(AttendanceRecord).filter(AttendanceRecord.session_id == session_id).all()}
    pending = {r.student_id: r for r in db.query(PendingAttendance).filter(PendingAttendance.session_id == session_id).all()}
    return csv_response("attendance.csv", ["name", "email", "registration_number", "cohort", "status", "source", "submitted_at", "verified_by"],
        [(r["student"].full_name, r["student"].email, r["code"], r["cohort"],
          "Present" if r["student"].id in records else (pending[r["student"].id].status if r["student"].id in pending else "Absent"),
          "Admin/company laptop" if r["student"].id in pending else "Attendance link",
          records[r["student"].id].checked_in_at if r["student"].id in records else pending[r["student"].id].submitted_at if r["student"].id in pending else "",
          (db.get(User, pending[r["student"].id].verified_by).name if r["student"].id in pending and pending[r["student"].id].verified_by else "")) for r in attendance_roster(db, session)])
